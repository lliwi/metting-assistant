"""Exportación de una grabación: transcripción, resumen y acta en varios formatos."""

import html
import io
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any

LABELS: dict[str, dict[str, str]] = {
    "es": {
        "transcript": "Transcripción",
        "summary": "Resumen",
        "key_points": "Puntos clave",
        "minutes": "Acta de la reunión",
        "attendees": "Asistentes",
        "topics": "Temas tratados",
        "decisions": "Decisiones",
        "tasks": "Tareas",
        "owner": "Responsable",
        "due": "Fecha límite",
        "open_questions": "Preguntas abiertas",
        "none": "Ninguno.",
        "speaker": "Hablante",
        "date": "Fecha",
        "duration": "Duración",
        "ai_notice": "Contenido generado por IA: puede contener errores y requiere revisión humana.",
    },
    "ca": {
        "transcript": "Transcripció",
        "summary": "Resum",
        "key_points": "Punts clau",
        "minutes": "Acta de la reunió",
        "attendees": "Assistents",
        "topics": "Temes tractats",
        "decisions": "Decisions",
        "tasks": "Tasques",
        "owner": "Responsable",
        "due": "Data límit",
        "open_questions": "Preguntes obertes",
        "none": "Cap.",
        "speaker": "Parlant",
        "date": "Data",
        "duration": "Durada",
        "ai_notice": "Contingut generat per IA: pot contenir errors i requereix revisió humana.",
    },
}


@dataclass(frozen=True)
class Segment:
    start_ms: int
    end_ms: int
    speaker: str | None
    text: str


@dataclass
class ExportContext:
    title: str
    created_at: datetime
    duration_ms: int
    locale: str
    segments: list[Segment] = field(default_factory=list)
    speaker_names: dict[str, str] = field(default_factory=dict)
    summary: dict[str, Any] | None = None
    minutes: dict[str, Any] | None = None

    @property
    def labels(self) -> dict[str, str]:
        return LABELS.get(self.locale, LABELS["es"])

    def speaker(self, label: str | None) -> str | None:
        if not label:
            return None
        return (
            self.speaker_names.get(label) or f"{self.labels['speaker']} {label.removeprefix('H')}"
        )


def clock(ms: int, *, fraction: str | None = None) -> str:
    """mm:ss (o h:mm:ss). Con `fraction` ("," o "."), formato de subtítulos hh:mm:ss,mmm."""
    if fraction:
        h, rest = divmod(ms, 3_600_000)
        m, rest = divmod(rest, 60_000)
        s, millis = divmod(rest, 1000)
        return f"{h:02d}:{m:02d}:{s:02d}{fraction}{millis:03d}"
    s = ms // 1000
    h, m, s = s // 3600, (s % 3600) // 60, s % 60
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def _refs(item: dict[str, Any]) -> str:
    refs = item.get("segmentos") or []
    return ", ".join(clock(r["start_ms"]) for r in refs if isinstance(r, dict))


# --- Estructura común (bloques) ------------------------------------------------------------
# Cada formato rico (Markdown, HTML/PDF, DOCX) recorre la misma lista de bloques.


@dataclass(frozen=True)
class Block:
    kind: str  # h1 | h2 | h3 | p | li | meta | notice
    text: str
    ref: str = ""  # marcas de tiempo de las citas, si las hay


def blocks(ctx: ExportContext) -> list[Block]:
    lb = ctx.labels
    out = [
        Block("h1", ctx.title),
        Block(
            "meta",
            f"{lb['date']}: {ctx.created_at:%d/%m/%Y %H:%M} · {lb['duration']}: {clock(ctx.duration_ms)}",
        ),
    ]
    if ctx.summary:
        out.append(Block("h2", lb["summary"]))
        out += [Block("p", p) for p in str(ctx.summary.get("resumen", "")).split("\n") if p.strip()]
        out.append(Block("h3", lb["key_points"]))
        out += _items(ctx.summary.get("puntos_clave"), lambda i: i.get("texto", ""), lb)
        out.append(Block("notice", lb["ai_notice"]))
    if ctx.minutes:
        m = ctx.minutes
        out.append(Block("h2", lb["minutes"]))
        if m.get("titulo"):
            out.append(Block("p", str(m["titulo"])))
        out.append(Block("h3", lb["attendees"]))
        out.append(Block("p", ", ".join(m.get("asistentes") or []) or lb["none"]))
        out.append(Block("h3", lb["topics"]))
        for topic in m.get("temas") or []:
            out.append(
                Block("p", f"{topic.get('titulo', '')}: {topic.get('resumen', '')}", _refs(topic))
            )
        out.append(Block("h3", lb["decisions"]))
        out += _items(m.get("decisiones"), lambda i: i.get("texto", ""), lb)
        out.append(Block("h3", lb["tasks"]))

        def task(i: dict[str, Any]) -> str:
            extra = [
                f"{lb['owner']}: {i['responsable']}" if i.get("responsable") else "",
                f"{lb['due']}: {i['fecha_limite']}" if i.get("fecha_limite") else "",
            ]
            details = " · ".join(e for e in extra if e)
            return f"{i.get('descripcion', '')}{f' ({details})' if details else ''}"

        out += _items(m.get("tareas"), task, lb)
        out.append(Block("h3", lb["open_questions"]))
        out += _items(m.get("preguntas_abiertas"), lambda i: i.get("texto", ""), lb)
        out.append(Block("notice", lb["ai_notice"]))
    if ctx.segments:
        out.append(Block("h2", lb["transcript"]))
        for seg in ctx.segments:
            who = ctx.speaker(seg.speaker)
            out.append(Block("p", f"{who + ': ' if who else ''}{seg.text}", clock(seg.start_ms)))
    return out


def _items(items: Any, render: Any, lb: dict[str, str]) -> list[Block]:
    if not items:
        return [Block("p", lb["none"])]
    return [Block("li", str(render(i)), _refs(i)) for i in items if isinstance(i, dict)]


# --- Formatos ------------------------------------------------------------------------------


def to_markdown(ctx: ExportContext) -> bytes:
    lines: list[str] = []
    in_transcript = False
    for b in blocks(ctx):
        ref = f" _[{b.ref}]_" if b.ref else ""
        match b.kind:
            case "h1":
                lines += [f"# {b.text}", ""]
            case "h2":
                in_transcript = b.text == ctx.labels["transcript"]
                lines += ["", f"## {b.text}", ""]
            case "h3":
                lines += ["", f"### {b.text}", ""]
            case "li":
                lines.append(f"- {b.text}{ref}")
            case "meta" | "notice":
                lines += ["", f"> {b.text}", ""]
            case _:
                # En la transcripción la marca va delante, como en la web.
                if in_transcript and b.ref:
                    lines += [f"**[{b.ref}]** {b.text}", ""]
                else:
                    lines += [f"{b.text}{ref}", ""]
    return ("\n".join(lines).strip() + "\n").encode()


def to_text(ctx: ExportContext) -> bytes:
    lines = [ctx.title, ""]
    for seg in ctx.segments:
        who = ctx.speaker(seg.speaker)
        lines.append(f"[{clock(seg.start_ms)}] {who + ': ' if who else ''}{seg.text}")
    return ("\n".join(lines) + "\n").encode()


def _subtitles(ctx: ExportContext, fraction: str) -> list[str]:
    cues = []
    for i, seg in enumerate(ctx.segments, start=1):
        who = ctx.speaker(seg.speaker)
        text = f"{who}: {seg.text}" if who else seg.text
        timing = (
            f"{clock(seg.start_ms, fraction=fraction)} --> {clock(seg.end_ms, fraction=fraction)}"
        )
        cues.append(f"{i}\n{timing}\n{text}\n")
    return cues


def to_srt(ctx: ExportContext) -> bytes:
    return "\n".join(_subtitles(ctx, ",")).encode()


def to_vtt(ctx: ExportContext) -> bytes:
    return ("WEBVTT\n\n" + "\n".join(_subtitles(ctx, "."))).encode()


_CSS = """
@page { size: A4; margin: 2cm; @bottom-right { content: counter(page); font-size: 9pt; color: #777; } }
body { font-family: "DejaVu Sans", sans-serif; font-size: 10.5pt; line-height: 1.45; color: #1b1d21; }
h1 { font-size: 20pt; margin: 0 0 4pt; }
h2 { font-size: 14pt; margin: 18pt 0 6pt; border-bottom: 1px solid #ccc; padding-bottom: 2pt; }
h3 { font-size: 11.5pt; margin: 12pt 0 4pt; }
p { margin: 0 0 6pt; }
ul { margin: 0 0 6pt; padding-left: 16pt; }
.meta { color: #555; }
.notice { color: #5c4300; background: #fff6e0; padding: 4pt 6pt; font-size: 9pt; }
.ref { color: #777; font-size: 8.5pt; }
.ts { color: #777; font-variant-numeric: tabular-nums; }
"""


def to_html(ctx: ExportContext) -> str:
    parts = [
        f'<!doctype html><html lang="{ctx.locale}"><meta charset="utf-8"><style>{_CSS}</style><body>'
    ]
    in_list = False
    in_transcript = False
    for b in blocks(ctx):
        text = html.escape(b.text)
        ref = f' <span class="ref">[{html.escape(b.ref)}]</span>' if b.ref else ""
        if b.kind != "li" and in_list:
            parts.append("</ul>")
            in_list = False
        match b.kind:
            case "h1":
                parts.append(f"<h1>{text}</h1>")
            case "h2":
                in_transcript = b.text == ctx.labels["transcript"]
                parts.append(f"<h2>{text}</h2>")
            case "h3":
                parts.append(f"<h3>{text}</h3>")
            case "li":
                if not in_list:
                    parts.append("<ul>")
                    in_list = True
                parts.append(f"<li>{text}{ref}</li>")
            case "meta":
                parts.append(f'<p class="meta">{text}</p>')
            case "notice":
                parts.append(f'<p class="notice">{text}</p>')
            case _:
                if in_transcript and b.ref:
                    parts.append(f'<p><span class="ts">[{html.escape(b.ref)}]</span> {text}</p>')
                else:
                    parts.append(f"<p>{text}{ref}</p>")
    if in_list:
        parts.append("</ul>")
    parts.append("</body></html>")
    return "".join(parts)


def to_pdf(ctx: ExportContext) -> bytes:
    from weasyprint import HTML

    # Sin URL base ni recursos externos: el HTML se genera aquí y no carga nada de fuera.
    pdf: bytes = HTML(string=to_html(ctx)).write_pdf()
    return pdf


def to_docx(ctx: ExportContext) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    in_transcript = False
    for b in blocks(ctx):
        match b.kind:
            case "h1":
                doc.add_heading(b.text, level=0)
            case "h2":
                in_transcript = b.text == ctx.labels["transcript"]
                doc.add_heading(b.text, level=1)
            case "h3":
                doc.add_heading(b.text, level=2)
            case "li":
                p = doc.add_paragraph(b.text, style="List Bullet")
                if b.ref:
                    run = p.add_run(f" [{b.ref}]")
                    run.font.size = Pt(8)
                    run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
            case "meta" | "notice":
                run = doc.add_paragraph().add_run(b.text)
                run.italic = True
                run.font.size = Pt(9)
            case _:
                p = doc.add_paragraph()
                if in_transcript and b.ref:
                    stamp = p.add_run(f"[{b.ref}] ")
                    stamp.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
                    p.add_run(b.text)
                else:
                    p.add_run(b.text)
                    if b.ref:
                        run = p.add_run(f" [{b.ref}]")
                        run.font.size = Pt(8)
                        run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


FORMATS = {
    "md": (to_markdown, "text/markdown; charset=utf-8"),
    "txt": (to_text, "text/plain; charset=utf-8"),
    "srt": (to_srt, "application/x-subrip; charset=utf-8"),
    "vtt": (to_vtt, "text/vtt; charset=utf-8"),
    "pdf": (to_pdf, "application/pdf"),
    "docx": (
        to_docx,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ),
}
